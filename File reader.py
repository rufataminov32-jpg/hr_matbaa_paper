import fitz          # PyMuPDF  — PDF
import docx           # python-docx — Word
import openpyxl       # openpyxl  — Excel


def pdf_dan_matn(file_path: str) -> list[tuple[str, str]]:
    """PDF fayldan (savol, javob) juftlarini qaytaradi.
    Savol = paragrafning birinchi 60 belgisi, javob = to'liq paragraf."""
    doc = fitz.open(file_path)
    juftlar = []
    for page in doc:
        matn = page.get_text("text")
        paragraflar = [p.strip() for p in matn.split("\n\n") if len(p.strip()) > 10]
        for p in paragraflar:
            savol = p[:60] + ("..." if len(p) > 60 else "")
            juftlar.append((savol, p))
    doc.close()
    return juftlar


def word_dan_matn(file_path: str) -> list[tuple[str, str]]:
    """Word (.docx) fayldan paragraflarni o'qiydi."""
    document = docx.Document(file_path)
    juftlar = []
    for para in document.paragraphs:
        matn = para.text.strip()
        if len(matn) > 10:
            savol = matn[:60] + ("..." if len(matn) > 60 else "")
            juftlar.append((savol, matn))
    return juftlar


def excel_dan_matn(file_path: str) -> list[tuple[str, str]]:
    """Excel (.xlsx) fayldan A ustun=savol, B ustun=javob o'qiydi.
    Birinchi qator sarlavha sifatida o'tkazib yuboriladi."""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    juftlar = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and len(row) >= 2 and row[0] and row[1]:
            savol = str(row[0]).strip()
            javob = str(row[1]).strip()
            if savol and javob:
                juftlar.append((savol, javob))
    wb.close()
    return juftlar
